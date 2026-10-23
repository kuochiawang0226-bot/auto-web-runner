# -*- coding: utf-8 -*-
"""GitHub Actions 單列執行版。請配合 .github/workflows/run-row.yml 與 README 使用。
每次工作只查一次 IP；重複時由下一個獨立工作重試，不使用 Tor。
原有網頁定位及操作邏輯保留。操作完成不等同網站已確認受理。
"""
import os
import time

_raw_url = os.environ.get("TARGET_URL", "")
TARGET_URL = _raw_url.strip()

print("網址設定檢查：", {
    "有讀到內容": bool(_raw_url),
    "前後有空白或換行": _raw_url != TARGET_URL,
    "以https開頭": TARGET_URL.startswith("https://"),
    "包含預設提示文字": "請填入" in TARGET_URL,
})
EXCEL_PATH = os.environ.get("EXCEL_PATH", "data.xlsx")
SHEET_NAME = os.environ.get("SHEET_NAME") or None
DROPDOWN1_OPTION = "中部汽車"
SUBMIT_TEXT = "送出"
WAIT_AFTER_SUBMIT = 5
DROPDOWN1_SELECTOR = ""
DROPDOWN2_SELECTOR = ""
FIELD_SELECTOR = ""
FIELD_PLACEHOLDER = "請填入文字"

CUSTOM_DROPDOWN_CSS = (
    '[role="listbox"][aria-expanded], '
    "[role=combobox], [aria-haspopup=listbox], .el-select, .ant-select, "
    ".select2-selection, .v-select, .dropdown-toggle"
)


def _select_native(sel, label, wait=15):
    """一般下拉選單：等選項出現（第二個選單常要等第一個選完才載入），再選取。"""
    deadline = time.time() + wait
    while time.time() < deadline:
        texts = [t.strip() for t in sel.locator("option").all_inner_texts()]
        if label in texts:
            sel.select_option(index=texts.index(label), force=True)
            return True
        partial = [k for k, t in enumerate(texts) if label in t]
        if partial:
            sel.select_option(index=partial[0], force=True)
            return True
        time.sleep(0.5)
    return False


def _select_custom(page, nth, label):
    """網站自製選單：點開第 nth 個選單，再點選項文字。"""
    page.locator(CUSTOM_DROPDOWN_CSS).nth(nth).click()
    time.sleep(0.8)
    opt = page.get_by_role("option", name=label)
    if opt.count() == 0:
        opt = page.get_by_text(label, exact=True)
    opt.last.click()


def choose(page, nth, label, selector=""):
    """選第 nth 個下拉選單（0＝第一個、1＝第二個）的指定項目。"""
    if selector:
        if not _select_native(page.locator(selector).first, label):
            raise RuntimeError(f"選單 {selector} 裡找不到「{label}」")
        return
    selects = page.locator("select")
    if selects.count() > nth and _select_native(selects.nth(nth), label):
        time.sleep(1)   # 讓網站有時間載入下一個選單的選項
        return
    _select_custom(page, nth, label)
    time.sleep(1)


TEXT_INPUT_CSS = (
    "input:not([type]), input[type=text], input[type=tel], input[type=email], "
    "input[type=number], input[type=search], textarea, [contenteditable=true]"
)


def _find_text_box(page, wait=15):
    """找出真正要打字的欄位：排除唯讀、停用、以及下拉選單內部的輸入框。"""
    deadline = time.time() + wait
    while time.time() < deadline:
        boxes = page.locator(TEXT_INPUT_CSS)
        for k in range(boxes.count()):
            box = boxes.nth(k)
            try:
                if not box.is_visible() or not box.is_editable():
                    continue
                inside_dropdown = box.evaluate(
                    """(el, css) => !!el.closest(css) || el.getAttribute('role') === 'combobox'
                         || el.hasAttribute('aria-autocomplete')
                         || el.getAttribute('aria-hidden') === 'true'
                         || el.getAttribute('tabindex') === '-1'
                         || /nativeInput/i.test(el.className)""",
                    CUSTOM_DROPDOWN_CSS + ", .select2-container, [role=listbox]",
                )
                if not inside_dropdown:
                    return box
            except Exception:
                continue
        time.sleep(0.5)
    raise RuntimeError("找不到可以輸入文字的欄位")


def fill_field(page, value):
    """模擬真人點欄位、清空、逐字輸入，讓網站確實收到內容。"""
    if FIELD_SELECTOR:
        box = page.locator(FIELD_SELECTOR).first
    elif FIELD_PLACEHOLDER and page.get_by_placeholder(FIELD_PLACEHOLDER).count() > 0:
        box = page.get_by_placeholder(FIELD_PLACEHOLDER).first
    else:
        box = _find_text_box(page)
    box.scroll_into_view_if_needed()
    box.click()
    box.press("Control+A")
    box.press("Delete")
    box.press_sequentially(value, delay=60)
    box.dispatch_event("change")
    box.press("Tab")          # 離開欄位，觸發網站的檢查
    time.sleep(0.5)
    try:
        got = box.input_value()
    except Exception:
        got = box.inner_text()
    if got.strip() != value.strip():
        raise RuntimeError(f"輸入欄位內容不符（預期「{value}」，實際「{got}」）")


def report_page(page, row_no):
    """失敗時印出網頁上的選單資訊並截圖，方便判斷原因。"""
    try:
        page.screenshot(path=f"錯誤畫面_第{row_no}列.png", full_page=True)
        selects = page.locator("select")
        print(f"  網頁上找到 {selects.count()} 個標準下拉選單：")
        for k in range(selects.count()):
            opts = [t.strip() for t in selects.nth(k).locator("option").all_inner_texts()][:8]
            print(f"   第 {k + 1} 個選項：{opts}")
        print(f"  自製選單數量：{page.locator(CUSTOM_DROPDOWN_CSS).count()}")
        boxes = page.locator(TEXT_INPUT_CSS)
        print(f"  找到 {boxes.count()} 個文字欄位：")
        for k in range(boxes.count()):
            b = boxes.nth(k)
            info = b.evaluate("e => [e.tagName, e.id, e.name, e.placeholder, e.readOnly, e.className].join(' | ')")
            print(f"   第 {k + 1} 個：{info}，看得到={b.is_visible()}")
        print(f"  已存截圖：錯誤畫面_第{row_no}列.png")
    except Exception:
        pass


def click_submit(page):
    """按下送出按鈕。"""
    btn = page.get_by_role("button", name=SUBMIT_TEXT)
    if btn.count() > 0:
        btn.first.click()
    else:
        page.locator(f"input[type=submit][value='{SUBMIT_TEXT}'], text={SUBMIT_TEXT}").first.click()


# ---------------- GitHub 工作排程與跨次紀錄 ----------------
import base64
import csv
import hashlib
import ipaddress
import json
import re
import urllib.request
import urllib.error
from pathlib import Path
from datetime import datetime, timezone, timedelta
from zoneinfo import ZoneInfo
import secrets

STATE_PATH = '.automation/state.json'
WORKFLOW_FILE = 'run-row.yml'
STOP_PATH = '.automation/STOP'


class StopRequested(Exception):
    pass


class SchedulePause(StopRequested):
    pass


TAIPEI = ZoneInfo('Asia/Taipei')


def local_now():
    return datetime.now(TAIPEI)


def day_plan(batch, clock):
    day = clock.date().isoformat()
    if batch.get('daily', {}).get('date') != day:
        base = clock.replace(hour=0, minute=0, second=0, microsecond=0)
        batch['daily'] = {
            'date': day,
            'start_at': (base + timedelta(hours=8, minutes=30, seconds=secrets.randbelow(1800))).isoformat(),
            'stop_at': (base + timedelta(hours=20, minutes=30, seconds=secrets.randbelow(1801))).isoformat(),
            'started': False, 'submitted': False}
    return batch['daily']


def check_daily_window(batch, clock):
    plan = batch.get('daily', {})
    if plan.get('date') != clock.date().isoformat() or not plan.get('started'):
        raise SchedulePause('等待今日排程啟動。')
    if not datetime.fromisoformat(plan['start_at']) <= clock < datetime.fromisoformat(plan['stop_at']):
        raise SchedulePause('已到休息時間，保留列號與輪次。')
    if not plan.get('submitted') and clock >= clock.replace(hour=9, minute=0, second=0, microsecond=0):
        raise SchedulePause('今天尚未送出且已超過 09:00，等待明日。')


def advance_round(batch):
    if batch['status'] not in ('complete', 'limit_reached'):
        return False
    if batch['round'] >= batch['max_rounds']:
        batch['status'] = 'all_rounds_complete'
        return False
    batch['round'] += 1
    batch['index'] = 0
    batch['runs'] = 0
    batch['retries'] = 0
    batch['status'] = 'ready'
    return True


def now():
    return datetime.now(timezone.utc).isoformat()


def valid_ip(value):
    ip = ipaddress.ip_address(value.strip())
    if not ip.is_global:
        raise ValueError('查詢結果不是公開 IP')
    return str(ip)


def legacy_ips(path):
    """匯入舊 CSV 所有有效 IPv4/IPv6，包含失敗列與附註。"""
    result = set()
    if Path(path).exists():
        with open(path, encoding='utf-8-sig', newline='') as f:
            for row in csv.reader(f):
                if len(row) < 5:
                    continue
                value = re.split(r'[（(\s]', row[4].strip(), maxsplit=1)[0]
                try:
                    result.add(valid_ip(value))
                except ValueError:
                    pass
    return result


class GitHubStore:
    def __init__(self):
        self.repo = os.environ['GITHUB_REPOSITORY']
        self.branch = os.environ['DEFAULT_BRANCH']
        self.token = os.environ['GH_TOKEN']
        self.api = os.environ.get('GITHUB_API_URL', 'https://api.github.com')
        self.sha = None

    def request(self, method, path, body=None):
        data = None if body is None else json.dumps(body).encode()
        req = urllib.request.Request(self.api + '/repos/' + self.repo + path,
            data=data, method=method, headers={
                'Authorization': 'Bearer ' + self.token,
                'Accept': 'application/vnd.github+json',
                'Content-Type': 'application/json',
                'User-Agent': 'single-row-automation'})
        with urllib.request.urlopen(req, timeout=45) as response:
            raw = response.read()
            return json.loads(raw) if raw else None

    def load(self):
        from urllib.parse import quote
        try:
            response = self.request('GET', '/contents/' + STATE_PATH + '?ref=' + quote(self.branch, safe=''))
        except urllib.error.HTTPError as e:
            if e.code != 404:
                raise
            return {'version': 1, 'used_ips': [], 'batches': {}}
        self.sha = response['sha']
        if not response.get('content'):
            response = self.request('GET', '/git/blobs/' + self.sha)
        state = json.loads(base64.b64decode(response['content']))
        if state.get('version') != 1:
            raise RuntimeError('不支援的狀態檔版本')
        return state

    def save(self, state):
        raw = json.dumps(state, ensure_ascii=False, indent=2).encode()
        body = {'message': 'Save automation progress', 'branch': self.branch,
                'content': base64.b64encode(raw).decode()}
        if self.sha:
            body['sha'] = self.sha
        # 使用 SHA 防止覆蓋其他工作的更新；失敗就停止，不自動覆寫。
        response = self.request('PUT', '/contents/' + STATE_PATH, body)
        self.sha = response['content']['sha']

    def check_stop(self):
        from urllib.parse import quote
        try:
            self.request('GET', '/contents/' + STOP_PATH + '?ref=' + quote(self.branch, safe=''))
        except urllib.error.HTTPError as exc:
            if exc.code == 404:
                return
            raise  # 無法確認停止旗標時，停止操作。
        raise StopRequested('偵測到 .automation/STOP，停止後續操作。')

    def dispatch(self, inputs):
        self.request('POST', '/actions/workflows/' + WORKFLOW_FILE + '/dispatches',
                     {'ref': self.branch, 'inputs': inputs})


def read_data(start, end):
    from openpyxl import load_workbook
    wb = load_workbook(EXCEL_PATH, data_only=True, read_only=True)
    try:
        ws = wb[SHEET_NAME] if SHEET_NAME else wb.worksheets[0]
        result = []
        for r, (a, b) in enumerate(ws.iter_rows(min_row=start, max_row=end or ws.max_row,
                                               max_col=2, values_only=True), start):
            if a is None or not str(a).strip():
                break
            result.append((r, str(a).strip(), '' if b is None else str(b).strip()))
        return result
    finally:
        wb.close()


def lookup_ip(page):
    response = page.goto('https://api.ipify.org', timeout=60000)
    if response is None or not response.ok:
        raise RuntimeError('IP 查詢服務回應失敗')
    return valid_ip(page.inner_text('body'))


def apply_row(page, row, check_stop=lambda: None):
    row_no, a, b = row
    page.goto(TARGET_URL, timeout=90000)
    page.wait_for_load_state('networkidle', timeout=60000)
    choose(page, 0, DROPDOWN1_OPTION, DROPDOWN1_SELECTOR)
    choose(page, 1, a, DROPDOWN2_SELECTOR)
    fill_field(page, b)
    check_stop()
    click_submit(page)
    time.sleep(WAIT_AFTER_SUBMIT)


def process_attempt(state, batch, rows, get_ip, operate, save):
    """一次只嘗試一個 IP；回傳 True 才接續新工作。可離線測試。"""
    if batch['status'] != 'ready':
        return False
    if batch['index'] >= len(rows):
        batch['status'] = 'complete'
        save(state)
        return False
    if batch['runs'] >= batch['max_runs']:
        batch['status'] = 'limit_reached'
        save(state)
        return False
    batch['runs'] += 1
    batch['step'] += 1
    row = rows[batch['index']]
    event = {'round': batch['round'], 'time': now(), 'row': row[0], 'run_id': os.environ.get('GITHUB_RUN_ID', ''),
             'result': 'checking', 'ip': ''}
    batch['events'].append(event)
    try:
        ip = valid_ip(get_ip())
    except Exception as exc:
        event['result'] = 'ip_lookup_failed'
        event['error'] = str(exc)
        batch['status'] = 'ip_lookup_failed'
        save(state)
        return False
    event['ip'] = ip
    if ip in state['used_ips']:
        event['result'] = 'duplicate_ip'
        batch['retries'] += 1
        # max_retries=5：首次重複後最多另外啟動 5 次。
        if batch['retries'] > batch['max_retries']:
            batch['status'] = 'retry_limit'
        elif batch['runs'] >= batch['max_runs']:
            batch['status'] = 'limit_reached'
        save(state)
        return batch['status'] == 'ready'
    # 操作網站前先保留 IP 及 in_flight，當機後不會默默重送。
    state['used_ips'].append(ip)
    batch['status'] = 'in_flight'
    event['result'] = 'in_flight'
    save(state)
    try:
        operate(row)
    except SchedulePause:
        event['result'] = 'paused_before_submit'
        batch['status'] = 'waiting_window'
        # 尚未送出，歸還本輪預算，確保隔日仍從同一列續做。step 不回退。
        batch['runs'] -= 1
        save(state)
        return False
    except StopRequested:
        event['result'] = 'stopped_before_submit'
        batch['status'] = 'stopped'
        batch['runs'] -= 1
        save(state)
        return False
    except Exception as exc:
        event['result'] = 'operation_uncertain'
        event['error'] = str(exc)
        batch['status'] = 'operation_uncertain'
        save(state)
        return False
    event['result'] = '操作完成（沿用原程式判定）'
    batch['index'] += 1
    batch['total_processed'] += 1
    batch['retries'] = 0
    batch['status'] = 'complete' if batch['index'] == len(rows) else 'ready'
    if batch['status'] == 'ready' and batch['runs'] >= batch['max_runs']:
        batch['status'] = 'limit_reached'
    save(state)
    return batch['status'] == 'ready'



def write_result(batch):
    Path('本次結果.json').write_text(json.dumps(batch, ensure_ascii=False, indent=2), encoding='utf-8')
    print('輪次：', batch['round'], '/', batch['max_rounds'],
          '本輪已處理：', batch['index'], '累計：', batch['total_processed'], '狀態：', batch['status'])
    if batch.get('daily'):
        print('今日台灣時間：', batch['daily']['start_at'], '～', batch['daily']['stop_at'])


def inputs_for(batch_id, batch):
    return {'mode': 'execute', 'batch_id': batch_id,
            'start_row': str(batch['start']), 'end_row': str(batch['end']),
            'max_retries': str(batch['max_retries']), 'max_runs': str(batch['max_runs']),
            'max_rounds': str(batch['max_rounds']), 'expected_step': str(batch['step']),
            'expected_date': batch['daily']['date']}


def start_day(batch, clock):
    plan = day_plan(batch, clock)
    if clock >= clock.replace(hour=9, minute=0, second=0, microsecond=0):
        return False
    if clock < datetime.fromisoformat(plan['start_at']):
        return False
    plan['started'] = True
    if batch['status'] == 'waiting_window':
        batch['status'] = 'ready'
    return True


def main():
    if os.environ.get('GITHUB_ACTIONS') != 'true':
        raise RuntimeError('請透過 GitHub Actions 執行。')
    scheduled = os.environ.get('GITHUB_EVENT_NAME') == 'schedule'
    mode = 'schedule' if scheduled else os.environ.get('MODE', 'check_ip')
    if mode not in ('check_ip', 'execute', 'schedule'):
        raise ValueError('未知模式')
    store = GitHubStore()
    state = store.load()
    state['used_ips'] = sorted(set(state['used_ips']) | legacy_ips('送出紀錄.csv'))
    if mode == 'check_ip':
        from playwright.sync_api import sync_playwright
        with sync_playwright() as pw:
            browser = pw.chromium.launch(headless=True)
            try:
                ip = lookup_ip(browser.new_page())
                print('出口 IP：', ip, '歷史重複：', ip in state['used_ips'])
                print('僅查 IP，不登記排程、不操作網站。')
            finally:
                browser.close()
        return
    if '請填入' in TARGET_URL or not TARGET_URL.startswith('https://'):
        raise ValueError('請先設定 TARGET_URL Secret（HTTPS 網址）。')
    expected = os.environ.get('EXPECTED_STEP', '').strip() if not scheduled else ''
    continuation = bool(expected)
    registration = not scheduled and not continuation
    if registration:
        batch_id = os.environ.get('BATCH_ID', '').strip()
        if not re.fullmatch(r'[A-Za-z0-9_-]{1,64}', batch_id):
            raise ValueError('批次名稱只能含英數字、底線、連字號，1–64 字。')
        start = int(os.environ.get('START_ROW') or '2')
        end = int(os.environ.get('END_ROW') or '0')
        retries = int(os.environ.get('MAX_RETRIES') or '5')
        max_runs = int(os.environ.get('MAX_RUNS') or '100')
        rounds = int(os.environ.get('MAX_ROUNDS') or '1')
        if start < 2 or (end and end < start) or not 0 <= retries <= 20 or not 1 <= max_runs <= 500 or not 1 <= rounds <= 100:
            raise ValueError('列號／上限／輪數設定不正確。')
        rows = read_data(start, end)
        if not rows:
            raise ValueError('選定範圍沒有資料。')
        config = dict(start=start, end=end, max_retries=retries, max_runs=max_runs, max_rounds=rounds,
                      fingerprint=hashlib.sha256(json.dumps([TARGET_URL, SHEET_NAME, rows], ensure_ascii=False).encode()).hexdigest())
        batch = state['batches'].get(batch_id)
        if batch is None:
            batch = dict(config, round=1, total_processed=0, index=0, runs=0, retries=0, step=0, status='ready', events=[])
            state['batches'][batch_id] = batch
        elif any(batch.get(k) != v for k, v in config.items()):
            raise RuntimeError('批次設定／資料不同。延續循環版請沿用原參數；舊版無輪數批次不自動轉換。')
        active = state.get('daily_schedule', {})
        if active.get('enabled') and active.get('batch_id') != batch_id:
            raise RuntimeError('已有其他每日排程。先將 state.json 的 daily_schedule.enabled 改為 false，再登記新批次。')
        store.check_stop()
        if batch['status'] == 'stopped':
            batch['status'] = 'ready'
        if batch['status'] not in ('ready', 'waiting_window', 'complete', 'limit_reached', 'all_rounds_complete'):
            raise RuntimeError('批次需人工核對：' + batch['status'])
        state['daily_schedule'] = {'batch_id': batch_id, 'enabled': batch['status'] != 'all_rounds_complete'}
        store.save(state)
    else:
        active = state.get('daily_schedule', {})
        if not active.get('enabled'):
            print('沒有啟用中的每日排程。')
            return
        batch_id = active['batch_id']
        batch = state['batches'][batch_id]
        if continuation:
            if os.environ.get('BATCH_ID') != batch_id or int(expected) != batch['step']:
                print('過期或不同批次的接續工作，略過。')
                return
            if os.environ.get('EXPECTED_DATE') != local_now().date().isoformat():
                print('跨日的舊接續工作，略過；由每日排程喚醒。')
                return
        rows = read_data(batch['start'], batch['end'])
        fingerprint = hashlib.sha256(json.dumps([TARGET_URL, SHEET_NAME, rows], ensure_ascii=False).encode()).hexdigest()
        if not rows or fingerprint != batch['fingerprint']:
            raise RuntimeError('資料／網址／工作表已改變，停止自動續跑，請還原。')
    try:
        store.check_stop()
    except StopRequested:
        if batch['status'] in ('ready', 'waiting_window'):
            batch['status'] = 'stopped'
            store.save(state)
        write_result(batch)
        print('手動停止開關生效。')
        return
    if batch['status'] in ('complete', 'limit_reached'):
        advance_round(batch)
        store.save(state)
    if batch['status'] == 'all_rounds_complete':
        state['daily_schedule']['enabled'] = False
        store.save(state)
        write_result(batch)
        return
    if batch['status'] not in ('ready', 'waiting_window'):
        write_result(batch)
        print('不自動恢復需要人工核對或手動暫停的批次。')
        return
    clock = local_now()
    if not continuation:
        plan = day_plan(batch, clock)
        store.save(state)
        if scheduled and plan['started']:
            print('今日已啟動，不建立第二條接續工作。')
            return
        if not plan['started']:
            # 排程 08:15 先準備；短間隔檢查停止旗標，最多等到 09:00。
            if scheduled:
                while local_now() < datetime.fromisoformat(plan['start_at']):
                    try:
                        store.check_stop()
                    except StopRequested:
                        batch['status'] = 'stopped'
                        store.save(state)
                        write_result(batch)
                        return
                    time.sleep(min(20, max(0.1, (datetime.fromisoformat(plan['start_at']) - local_now()).total_seconds())))
            if not start_day(batch, local_now()):
                batch['status'] = 'waiting_window'
                store.save(state)
                write_result(batch)
                print('已登記每日排程，等待下一個允許的早晨時段。')
                return
            store.save(state)
    def guard():
        store.check_stop()
        check_daily_window(batch, local_now())
    try:
        guard()
    except SchedulePause:
        batch['status'] = 'waiting_window'
        store.save(state)
        write_result(batch)
        return
    except StopRequested:
        batch['status'] = 'stopped'
        store.save(state)
        write_result(batch)
        return
    if batch['status'] == 'waiting_window':
        batch['status'] = 'ready'
    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        browser = pw.chromium.launch(headless=True)
        try:
            page = browser.new_page()
            def operate(row):
                try:
                    apply_row(page, row, guard)
                    batch['daily']['submitted'] = True
                except StopRequested:
                    raise
                except Exception:
                    report_page(page, row[0])
                    raise
            more = process_attempt(state, batch, rows, lambda: lookup_ip(page), operate, store.save)
        finally:
            browser.close()
    if batch['status'] in ('complete', 'limit_reached'):
        more = advance_round(batch)
        store.save(state)
    if batch['status'] == 'all_rounds_complete':
        state['daily_schedule']['enabled'] = False
        store.save(state)
    write_result(batch)
    if more:
        time.sleep(15)
        try:
            guard()
        except SchedulePause:
            batch['status'] = 'waiting_window'
            store.save(state)
            write_result(batch)
            return
        except StopRequested:
            batch['status'] = 'stopped'
            store.save(state)
            write_result(batch)
            return
        store.dispatch(inputs_for(batch_id, batch))
        print('已要求下一個工作接續。')
    elif batch['status'] not in ('all_rounds_complete', 'stopped', 'waiting_window'):
        raise RuntimeError('批次已暫停：' + batch['status'])


if __name__ == '__main__':
    try:
        main()
    except StopRequested as exc:
        print(str(exc))
